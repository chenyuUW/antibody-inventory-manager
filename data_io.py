# data_io.py
# Excel read/write functions for Antibody Inventory Manager

import os
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from config import (
    INVENTORY_FILE,
    INVENTORY_SHEET,
    LOG_SHEET,
    INVENTORY_COLUMNS,
    LOG_COLUMNS,
)


def check_inventory_file_exists():
    """
    Check whether inventory Excel file exists.
    """
    if not os.path.exists(INVENTORY_FILE):
        raise FileNotFoundError(
            f"Inventory file '{INVENTORY_FILE}' was not found. "
            "Please make sure it is in the same folder as gate.py."
        )


def load_inventory():
    """
    Load inventory sheet into a pandas DataFrame.
    """
    check_inventory_file_exists()

    df = pd.read_excel(
        INVENTORY_FILE,
        sheet_name=INVENTORY_SHEET,
        dtype=str,
        engine="openpyxl",
    )

    # Make sure required columns exist
    missing_columns = [col for col in INVENTORY_COLUMNS if col not in df.columns]

    if missing_columns:
        raise ValueError(
            "The inventory sheet is missing required columns: "
            + ", ".join(missing_columns)
        )

    # Keep only expected columns in expected order
    df = df[INVENTORY_COLUMNS].copy()

    # Remaining should be numeric internally
    df["Remaining"] = pd.to_numeric(df["Remaining"], errors="coerce").fillna(0).astype(int)

    return df


def save_inventory(df):
    """
    Save inventory DataFrame back to the Excel file.

    This replaces the Inventory sheet while preserving other sheets when possible.
    """
    check_inventory_file_exists()

    try:
        # Make a backup before saving
        create_backup()

        with pd.ExcelWriter(
            INVENTORY_FILE,
            engine="openpyxl",
            mode="a",
            if_sheet_exists="replace",
        ) as writer:
            df.to_excel(writer, sheet_name=INVENTORY_SHEET, index=False)

    except PermissionError:
        raise PermissionError(
            f"Cannot save '{INVENTORY_FILE}'. "
            "The Excel file may be open in Excel or locked by another program. "
            "Please close the file and try again."
        )


def create_backup():
    """
    Create a timestamped backup copy before modifying the Excel file.
    """
    if not os.path.exists(INVENTORY_FILE):
        return

    backup_folder = "backups"

    if not os.path.exists(backup_folder):
        os.makedirs(backup_folder)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file = os.path.join(
        backup_folder,
        f"Antibody Inventory Master_backup_{timestamp}.xlsx"
    )

    with open(INVENTORY_FILE, "rb") as source:
        with open(backup_file, "wb") as target:
            target.write(source.read())


def ensure_log_sheet_exists():
    """
    Ensure the Log sheet exists.
    If not, create it with the required columns.
    """
    check_inventory_file_exists()

    workbook = load_workbook(INVENTORY_FILE)

    if LOG_SHEET not in workbook.sheetnames:
        worksheet = workbook.create_sheet(LOG_SHEET)

        for col_index, column_name in enumerate(LOG_COLUMNS, start=1):
            worksheet.cell(row=1, column=col_index, value=column_name)

        workbook.save(INVENTORY_FILE)


def append_log(log_record):
    """
    Append a log record to the Log sheet.

    log_record should be a dictionary with keys matching LOG_COLUMNS.
    """
    check_inventory_file_exists()
    ensure_log_sheet_exists()

    try:
        workbook = load_workbook(INVENTORY_FILE)
        worksheet = workbook[LOG_SHEET]

        next_row = worksheet.max_row + 1

        for col_index, column_name in enumerate(LOG_COLUMNS, start=1):
            worksheet.cell(
                row=next_row,
                column=col_index,
                value=log_record.get(column_name, "")
            )

        workbook.save(INVENTORY_FILE)

    except PermissionError:
        raise PermissionError(
            f"Cannot write log to '{INVENTORY_FILE}'. "
            "The Excel file may be open in Excel or locked by another program. "
            "Please close the file and try again."
        )

def get_latest_take_log(catalog, container_type=None):
    """
    Return the latest Take log matching the catalog
    and, when provided, the container type.

    Returns a dictionary or None.
    """
    check_inventory_file_exists()
    ensure_log_sheet_exists()

    log_df = pd.read_excel(
        INVENTORY_FILE,
        sheet_name=LOG_SHEET,
        dtype=str,
        engine="openpyxl",
    )

    if log_df.empty:
        return None

    required_columns = [
        "Time",
        "Action",
        "Catalog",
        "Container_Type",
        "Notes",
    ]

    for column in required_columns:
        if column not in log_df.columns:
            return None

    catalog_query = str(catalog).strip().upper()

    matched = log_df[
        (log_df["Action"].fillna("").str.strip().str.upper() == "TAKE")
        &
        (log_df["Catalog"].fillna("").str.strip().str.upper() == catalog_query)
    ].copy()

    if container_type:
        container_query = str(container_type).strip().upper()

        matched = matched[
            matched["Container_Type"]
            .fillna("")
            .str.strip()
            .str.upper()
            == container_query
        ]

    if matched.empty:
        return None

    latest = matched.iloc[-1]

    return {
        "Time": str(latest.get("Time", "")).strip(),
        "Notes": str(latest.get("Notes", "")).strip(),
    }