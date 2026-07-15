# gate.py
# Antibody Inventory Manager
# Main entrance file

import os
import sys
from datetime import datetime

from openpyxl import load_workbook

from config import (
    APP_NAME,
    VERSION,
    INVENTORY_FILE,
    INVENTORY_SHEET,
    COL_REMAINING,
)

# Load all main functions once when the program starts.
# This moves the waiting time to program startup, so opening each function is fast.
from take_antibody import take_antibody
from add_antibody import add_antibody
from search_inventory import search_inventory


# Cache inventory summary so the Excel file is not reread every time the menu appears.
_inventory_summary_cache = {
    "modified_time": None,
    "entry_count": 0,
    "total_remaining": 0,
    "last_updated": "Unknown",
    "error": None,
}


def clear_screen():
    """
    Clear terminal screen for Windows/macOS/Linux.
    """
    os.system("cls" if os.name == "nt" else "clear")


def pause():
    """
    Pause before returning to the main menu.
    """
    input("\nPress Enter to return to the main menu...")


def get_inventory_summary():
    """
    Return a short inventory summary.

    The Excel file is only reread when its modification time changes.
    This keeps menu refreshes fast while still updating the summary after
    inventory changes.
    """
    global _inventory_summary_cache

    if not os.path.exists(INVENTORY_FILE):
        return {
            "entry_count": 0,
            "total_remaining": 0,
            "last_updated": "File not found",
            "error": f"'{INVENTORY_FILE}' was not found",
        }

    modified_time = os.path.getmtime(INVENTORY_FILE)

    if _inventory_summary_cache["modified_time"] == modified_time:
        return _inventory_summary_cache

    try:
        workbook = load_workbook(
            INVENTORY_FILE,
            read_only=True,
            data_only=True,
        )

        if INVENTORY_SHEET not in workbook.sheetnames:
            workbook.close()
            raise ValueError(
                f"Worksheet '{INVENTORY_SHEET}' was not found."
            )

        worksheet = workbook[INVENTORY_SHEET]

        # Read the header row and locate the Remaining column.
        headers = {
            str(cell.value).strip(): index
            for index, cell in enumerate(worksheet[1], start=1)
            if cell.value is not None
        }

        remaining_column = headers.get(COL_REMAINING)

        entry_count = 0
        total_remaining = 0

        for row in worksheet.iter_rows(min_row=2, values_only=True):
            # Ignore completely empty rows.
            if not any(value not in (None, "") for value in row):
                continue

            entry_count += 1

            if remaining_column is not None:
                remaining_value = row[remaining_column - 1]

                try:
                    total_remaining += int(float(remaining_value or 0))
                except (TypeError, ValueError):
                    pass

        workbook.close()

        _inventory_summary_cache = {
            "modified_time": modified_time,
            "entry_count": entry_count,
            "total_remaining": total_remaining,
            "last_updated": datetime.fromtimestamp(modified_time).strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            "error": None,
        }

    except Exception as error:
        _inventory_summary_cache = {
            "modified_time": modified_time,
            "entry_count": 0,
            "total_remaining": 0,
            "last_updated": datetime.fromtimestamp(modified_time).strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            "error": str(error),
        }

    return _inventory_summary_cache


def print_header():
    """
    Print program header and current inventory summary.
    """
    summary = get_inventory_summary()

    print("=" * 70)
    print(f"{APP_NAME:^70}")
    print(f"{VERSION:^70}")
    print("=" * 70)
    print(f"Database file:         {INVENTORY_FILE}")
    print(f"Antibody entries:      {summary['entry_count']}")
    print(f"Total bottles/tubes:   {summary['total_remaining']}")
    print(f"Last inventory update: {summary['last_updated']}")
    print(f"Session started:       {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    if summary.get("error"):
        print(f"Inventory summary warning: {summary['error']}")

    print("=" * 70)


def print_menu():
    """
    Print main menu options.
    """
    print("\nMain Menu")
    print("-" * 70)
    print("1. Take antibody")
    print("2. Add / Register antibody")
    print("3. Search / Preview inventory")
    print("4. Exit")
    print("-" * 70)


def safe_run(function, display_name):
    """
    Safely run an already imported function.
    """
    try:
        print(f"\nOpening: {display_name}\n")
        function()

    except FileNotFoundError as error:
        print("\n[File error] Required file was not found.")
        print(f"Error details: {error}")

    except Exception as error:
        print("\n[Error] Something went wrong while running this function.")
        print(f"Error details: {error}")

def check_inventory_write_access():
    """
    Check whether the inventory file is available.

    Return True when the file can be accessed normally.
    Return False when it is missing or locked by Excel/another program.
    """
    if not os.path.exists(INVENTORY_FILE):
        print("\n" + "=" * 70)
        print("INVENTORY FILE NOT FOUND")
        print("=" * 70)
        print("The following file could not be found:")
        print(INVENTORY_FILE)
        print()
        print("Please make sure 'Antibody Inventory Master.xlsx'")
        print("is in the main project folder.")
        print("=" * 70)
        return False

    try:
        with open(INVENTORY_FILE, "r+b"):
            pass

        return True

    except PermissionError:
        print("\n" + "=" * 70)
        print("INVENTORY FILE IS CURRENTLY LOCKED")
        print("=" * 70)
        print("Antibody Inventory Master.xlsx appears to be open")
        print("in Microsoft Excel or another program.")
        print()
        print("Please close the Excel file before opening")
        print("Antibody Inventory Manager.")
        print("=" * 70)
        return False

    except OSError as error:
        print("\n" + "=" * 70)
        print("INVENTORY FILE ACCESS ERROR")
        print("=" * 70)
        print(f"Error details: {error}")
        print("=" * 70)
        return False

def main():
    """
    Main control loop.
    """
    while True:
        if not check_inventory_write_access():
            input("\nPress Enter to close...")
            sys.exit(1)

    session_started = datetime.now()

    while True:
        clear_screen()
        print_header()
        print_menu()

        choice = input("Select an option: ").strip()

        if choice == "1":
            clear_screen()
            print_header()
            safe_run(take_antibody, "Take antibody")
            pause()

        elif choice == "2":
            clear_screen()
            print_header()
            safe_run(add_antibody, "Add / Register antibody")
            pause()

        elif choice == "3":
            clear_screen()
            print_header()
            safe_run(search_inventory, "Search / Preview inventory")
            pause()

        elif choice == "4":
            clear_screen()
            print_header()
            print("\nGoodbye.")
            sys.exit(0)

        else:
            print("\nInvalid option. Please enter 1, 2, 3, or 4.")
            pause()

if __name__ == "__main__":
    main()
